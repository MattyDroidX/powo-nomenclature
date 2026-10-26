"""Spreadsheet in, review workbook and fill-only-empty import file out; the input file is never modified."""
from collections import Counter
from datetime import date
from pathlib import Path

from .match import NameQuery, match_name
from .resolve import FIELDS, resolve

# Column names of a Taxa/TSennen complete export (Sennen collection).
TAXA_COLUMNS = {'genus': 'porto_genus', 'species': 'porto_specific_epithet', 'author': 'porto_author',
                'infra_rank': 'porto_infra_rank', 'infra_epithet': 'porto_infra_epithet', 'infra_author': 'porto_infra_author'}
LABELS = {
    'en': {'results': 'Results', 'differences': 'Differences', 'source': 'Source', 'row': 'Row', 'name': 'Name as written',
           'method': 'Method', 'score': 'Score', 'review': 'Review', 'notes': 'Notes', 'wcvp_name': 'WCVP name as written',
           'accepted': 'Accepted name', 'field': 'Field', 'existing': 'Existing value', 'proposed': 'Proposed value',
           'yes': 'yes', 'no': 'no', 'version': 'WCVP version', 'citation': 'Citation', 'sha256': 'Source SHA-256',
           'run': 'Run date', 'rows': 'Rows with a name', 'filled': 'Fields filled in the import file'},
    'pt': {'results': 'Resultados', 'differences': 'Diferenças', 'source': 'Fonte', 'row': 'Linha', 'name': 'Nome como escrito',
           'method': 'Método', 'score': 'Pontuação', 'review': 'Rever', 'notes': 'Notas', 'wcvp_name': 'Nome WCVP correspondente',
           'accepted': 'Nome aceite', 'field': 'Campo', 'existing': 'Valor existente', 'proposed': 'Valor proposto',
           'yes': 'sim', 'no': 'não', 'version': 'Versão do WCVP', 'citation': 'Citação', 'sha256': 'SHA-256 da origem',
           'run': 'Data da execução', 'rows': 'Linhas com nome', 'filled': 'Campos preenchidos no ficheiro de importação'},
}


def _text(value):
    return '' if value is None else str(value).strip()


def _write(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'  # Names and notes are text, never Excel formulas.


def process_workbook(source, checklist, output_dir, *, sheet=None, columns=TAXA_COLUMNS, lang='en',
                     overwrite=False, include_review=False, checked_at=None):
    """Return (review_path, import_path or None, Counter of methods). Existing values are kept unless overwrite."""
    from openpyxl import Workbook, load_workbook
    labels = LABELS[lang]
    source, output_dir = Path(source), Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    book = load_workbook(source, read_only=True, data_only=True)
    try:
        rows = book[sheet] if sheet else book.worksheets[0]
        data = list(rows.iter_rows(values_only=True))
    finally:
        book.close()  # read-only workbooks keep the file open; Windows would refuse to move or delete it.
    header = [_text(h) for h in data[0]]
    missing = [c for c in (columns['genus'], columns['species']) if c not in header]
    if missing:
        raise ValueError(f'Missing name columns: {", ".join(missing)}')
    keyed = 'id' in header
    checked_at = checked_at or date.today()

    review_book = Workbook()
    results = review_book.active
    results.title = labels['results']
    _write(results, [labels['row'], *(['id'] if keyed else []), labels['name'], labels['method'], labels['score'],
                     labels['review'], labels['notes'], labels['wcvp_name'], labels['accepted'], *FIELDS])
    differences = review_book.create_sheet(labels['differences'])
    _write(differences, [labels['row'], *(['id'] if keyed else []), labels['name'], labels['field'], labels['existing'], labels['proposed']])
    import_rows, methods, filled = [], Counter(), 0
    for number, values in enumerate(data[1:], start=2):
        record = dict(zip(header, values))
        query = NameQuery(**{key: _text(record.get(column)) for key, column in columns.items()})
        if not query.genus or not query.species:
            continue
        match = match_name(checklist, query)
        result = resolve(checklist, match, checked_at=checked_at)
        review = match.review or result.review
        methods[match.method] += 1
        ident = [_text(record.get('id'))] if keyed else []
        original = match.original
        _write(results, [number, *ident, query.text(), match.method, match.score, labels['yes'] if review else labels['no'],
                         '; '.join(match.notes + result.notes),
                         f"{original['taxon_name']} {original['taxon_authors']} ({original['taxon_status']})".strip() if original else '',
                         result.accepted_name, *[result.values.get(f, '') for f in FIELDS]])
        proposed = {}
        for f in FIELDS:
            existing, value = _text(record.get(f)), _text(result.values.get(f))
            if existing and value and existing != value and f != 'powo_checked_at':
                _write(differences, [number, *ident, query.text(), f, existing, value])
            if value and (overwrite or not existing):
                proposed[f] = value
        if keyed and proposed and (include_review or not review):
            filled += len(proposed)
            import_rows.append([_text(record.get('id')), *([record.get('revision')] if 'revision' in header else []),
                                *[proposed.get(f, '' if overwrite else _text(record.get(f))) for f in FIELDS]])
    meta = review_book.create_sheet(labels['source'])
    for key, value in ((labels['version'], checklist.meta.get('wcvp_version')), (labels['citation'], checklist.meta.get('citation')),
                       (labels['sha256'], checklist.meta.get('source_sha256')), (labels['run'], checked_at.isoformat()),
                       (labels['rows'], sum(methods.values())), (labels['filled'], filled),
                       *((f"{labels['method']}: {m}", n) for m, n in methods.most_common())):
        _write(meta, [key, value])
    for sheet_ in (results, differences):
        sheet_.freeze_panes = 'B2'
        sheet_.auto_filter.ref = sheet_.dimensions
    review_path = output_dir / f'{source.stem}-powo-review.xlsx'
    review_book.save(review_path)
    import_path = None
    if keyed and import_rows:
        import_book = Workbook()
        sheet_ = import_book.active
        sheet_.title = 'Import'
        _write(sheet_, ['id', *(['revision'] if 'revision' in header else []), *FIELDS])
        for row in import_rows:
            _write(sheet_, row)
        import_path = output_dir / f'{source.stem}-powo-import.xlsx'
        import_book.save(import_path)
    return review_path, import_path, methods
