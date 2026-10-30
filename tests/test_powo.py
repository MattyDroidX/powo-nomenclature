"""Reference cases come from names a researcher resolved by hand on POWO (Sennen collection, Porto herbarium)."""
import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from datetime import date
from pathlib import Path

from powo_nomenclature import Checklist, NameQuery, build_index, match_name, resolve, verify
from powo_nomenclature.cli import main
from powo_nomenclature.normalize import author_similarity, clean_infra, epithet_key, epithet_variants, fold, genus_key

FIXTURE = Path(__file__).with_name('fixtures') / 'wcvp-mini.zip'


def url(ipni):
    return f'https://powo.science.kew.org/taxon/urn:lsid:ipni.org:names:{ipni}'


class NormalizeTests(unittest.TestCase):
    def test_ligatures_hybrid_signs_and_abbreviations(self):
        self.assertEqual(fold('Cratægus × Bonanovæ'), 'crataegus bonanovae')
        self.assertEqual(genus_key('Chærophyllum'), 'chaerophyllum')
        self.assertEqual(epithet_key('× vayredae'), 'vayredae')
        self.assertEqual(clean_infra('proles', 'S. iberica')[:2], ('proles', 'iberica'))
        self.assertEqual(clean_infra('fa.', 'berthulona')[:2], ('f.', 'berthulona'))

    def test_free_text_in_infraspecific_columns_is_reported(self):
        rank, epithet, notes = clean_infra('', "'= F. major x capreolata ej.")
        self.assertEqual((rank, epithet), ('', ''))
        self.assertIn('ignored', notes[0])
        self.assertEqual(clean_infra('var.', 'gracilentus inter fa. typ. et var. prostrat.')[1], 'gracilentus')

    def test_spelling_variants(self):
        self.assertEqual(epithet_variants('xqueri')[0], 'queri')
        self.assertIn('barcinonensis', epithet_variants('barcinonense'))
        self.assertNotIn('anthocarpa', epithet_variants('xanthocarpa'))

    def test_authors(self):
        self.assertEqual(author_similarity('Sen. et Pau', 'Sennen & Pau'), 1.0)
        self.assertEqual(author_similarity('L.', '(L.) Huds.'), 1.0)
        self.assertEqual(author_similarity('L.', 'Link'), 0.0)
        self.assertEqual(author_similarity('Cav. non L.', 'L.'), 0.0)


class ChecklistTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.index = Path(cls.tmp.name) / 'wcvp.sqlite'
        build_index(FIXTURE, cls.index)
        cls.checklist = Checklist(cls.index)

    @classmethod
    def tearDownClass(cls):
        cls.checklist.close()
        cls.tmp.cleanup()

    def run_name(self, *parts, author='', infra_author=''):
        query = NameQuery(*parts[:2], author, *(parts[2:4] if len(parts) > 2 else ('', '')), infra_author)
        match = match_name(self.checklist, query)
        return match, resolve(self.checklist, match, checked_at=date(2026, 10, 20))

    def test_index_records_its_source(self):
        self.assertEqual(self.checklist.meta['wcvp_version'], 'WCVP 2026-06-04')
        self.assertRegex(self.checklist.meta['source_sha256'], r'^[0-9a-f]{64}$')
        self.assertIn('CC BY 3.0', self.checklist.meta['citation'])
        with self.assertRaises(FileExistsError):
            build_index(FIXTURE, self.index)

    def test_synonym_status_and_link_describe_the_name_as_written(self):
        match, result = self.run_name('Micromeria', 'filiformis', 'var.', 'glandulosa')
        self.assertEqual((match.method, match.review), ('exact', False))
        values = result.values
        self.assertEqual(values['powo_url'], url('77283132-1'))
        self.assertEqual(values['powo_status'], 'This name is a synonym of Micromeria filiformis subsp. filiformis')
        self.assertEqual(values['powo_checked_at'], '20/10/2026')
        # Accepted infraspecific name: author goes to the infraspecific column (an autonym has none).
        self.assertEqual([values[f] for f in ('powo_family', 'powo_genus', 'powo_specific_epithet', 'powo_author',
                                               'powo_infra_rank', 'powo_infra_epithet', 'powo_infra_author')],
                         ['Lamiaceae', 'Micromeria', 'filiformis', '', 'subsp.', 'filiformis', ''])
        self.assertEqual(values['powo_native'], 'Baleares')
        self.assertIn('Micromeria filiformis var. glandulosa Sennen & Pau in Treb. Inst. Catalana Hist. Nat. 3: 193 (1917)',
                      values['powo_synonyms'].splitlines())

    def test_names_without_ipni_record_link_their_powo_page(self):
        self.assertEqual(self.run_name('Micromeria', 'filiformis', 'var.', 'rubrifolia')[1].values['powo_url'], url('550052-4'))

    def test_accepted_species_with_abbreviated_authors(self):
        match, result = self.run_name('Veronica', 'micrantha', author='Hoff. et LK.')
        self.assertEqual(result.values['powo_status'], 'This species is accepted')
        self.assertEqual((result.values['powo_url'], result.values['powo_author']), (url('812305-1'), 'Hoffmanns. & Link'))
        self.assertEqual(result.values['powo_native'], 'Portugal')

    def test_missing_form_falls_back_to_the_species_as_the_researcher_did(self):
        match, result = self.run_name('Xanthium', 'sallentii', 'f.', 'berthulona', author='Sen. Cad. et Pau')
        self.assertEqual((match.method, match.review), ('species-fallback', False))
        self.assertEqual(result.values['powo_url'], url('260886-1'))
        self.assertEqual(result.values['powo_status'], 'This name is a synonym of Xanthium × vayredae')
        self.assertEqual((result.values['powo_genus'], result.values['powo_specific_epithet'], result.values['powo_author']),
                         ('Xanthium', 'vayredae', 'Sennen'))
        # WCVP lists Spain as introduced for this hybrid: never reported as native.
        self.assertEqual(result.values['powo_native'], '')
        self.assertIn('introduced into: Spain', '; '.join(result.notes))

    def test_invalid_and_historical_names_always_need_review(self):
        match, result = self.run_name('Crataegus', 'brevispina', 'forma', 'subinermis', infra_author='Kze.')
        self.assertTrue(match.review)
        self.assertEqual(result.values['powo_status'], 'This name is a synonym of Crataegus monogyna')
        match, _ = self.run_name('Salvia', 'officinalis', 'proles', 'S. iberica')
        self.assertEqual((match.method, match.review), ('historical-rank', True))

    def test_spelling_cascade_is_flagged_for_review(self):
        for parts, method in ((('Erigeron', 'barcinonense'), 'orthographic'), (('Fumaria', 'xqueri'), 'orthographic'),
                              (('Calamintha', 'enriquesi'), 'fuzzy'), (('Cratægus', 'monogyna'), 'exact')):
            match, _ = self.run_name(*parts)
            self.assertEqual(match.method, method, parts)
            self.assertEqual(match.review, method != 'exact', parts)
        match, result = self.run_name('Viola', 'rolandi')
        self.assertEqual((match.method, match.review, result.values), ('none', True, {}))

    def test_homonyms_are_decided_by_author_or_reported(self):
        match, _ = self.run_name('Mentha', 'rotundifolia', author='L.')
        self.assertEqual(match.original['taxon_authors'], '(L.) Huds.')
        match, _ = self.run_name('Stellaria', 'media')
        self.assertTrue(match.review)
        self.assertIn('no decisive author', '; '.join(match.notes))

    def test_comparison_with_an_ai_filled_sheet_led_to_these_rules(self):
        # Spelling slips in the variety epithet are corrected within the species, never silently.
        match, _ = self.run_name('Cerastium', 'brachypetalum', 'var', 'eriocarlon')
        self.assertEqual((match.method, match.review, match.original['taxon_name']), ('infra-fuzzy', True, 'Cerastium brachypetalum var. eriocaulon'))
        # WCVP lists some names at historical ranks: the exact name wins, and it is still reviewed.
        match, _ = self.run_name('Stellaria', 'media', 'proles', 'catalaunica')
        self.assertEqual((match.method, match.review, match.original['taxon_name']), ('exact', True, 'Stellaria media proles catalaunica'))
        # The author decides between homonyms: «Lamk.» is Lamarck, the accepted name, not the synonym of Bartsia.
        match, result = self.run_name('Euphrasia', 'alpina', author='Lamk.')
        self.assertEqual((match.original['taxon_authors'], result.accepted_name), ('Lam.', 'Euphrasia alpina'))
        # Nothing found: the nearest spellings are offered to the reviewer, none is chosen.
        match, result = self.run_name('Pinguicula', 'eliasii')
        self.assertEqual((match.method, result.values), ('none', {}))
        self.assertIn('Pinguicula eliae', next(n for n in match.notes if n.startswith('closest WCVP names')))

    def test_species_author_for_infraspecific_names_is_an_option(self):
        match, _ = self.run_name('Micromeria', 'filiformis', 'var.', 'glandulosa')
        default = resolve(self.checklist, match).values['powo_author']
        species = resolve(self.checklist, match, species_author_for_infraspecific=True).values['powo_author']
        self.assertEqual(default, '')
        self.assertEqual(species, [r['taxon_authors'] for r in self.checklist.find('micromeria', 'filiformis') if r['taxon_status'] == 'Accepted'][0])

    def test_verify_flags_what_an_audit_should_catch(self):
        query = NameQuery('Carduus', 'nuriæ')
        unplaced = self.checklist.find('carduus', 'nuriae')[0]
        issues = verify(self.checklist, query, {'powo_genus': 'Carduus', 'powo_specific_epithet': 'nuriae',
                                                'powo_url': url(unplaced['powo_id'])})
        self.assertEqual([i.code for i in issues], ['unplaced_has_accepted'])
        issues = verify(self.checklist, NameQuery('Veronica', 'micrantha'),
                        {'powo_url': url('99999999-9'), 'powo_genus': 'Veronica', 'powo_specific_epithet': 'inventada'})
        self.assertEqual({i.code for i in issues}, {'unknown_page', 'accepted_not_accepted'})
        issues = verify(self.checklist, NameQuery('Micromeria', 'filiformis', infra_rank='var.', infra_epithet='glandulosa'),
                        {'powo_url': url('812305-1'), 'powo_genus': 'Micromeria', 'powo_specific_epithet': 'filiformis',
                         'powo_infra_rank': 'subsp.', 'powo_infra_epithet': 'filiformis'}, lang='pt')
        self.assertEqual({i.code for i in issues}, {'link_not_name_as_written', 'link_accepted_mismatch'})
        self.assertIn('Veronica micrantha', issues[0].message)
        # The tool's own result passes its own audit.
        match, result = self.run_name('Micromeria', 'filiformis', 'var.', 'glandulosa')
        self.assertEqual(verify(self.checklist, match.query, result.values), [])

    def test_indexes_from_another_version_are_refused_with_the_command_to_rebuild(self):
        import sqlite3
        old = Path(self.tmp.name) / 'old.sqlite'
        with sqlite3.connect(old) as db:
            db.execute('CREATE TABLE meta(key TEXT PRIMARY KEY, value TEXT)')
            db.execute("INSERT INTO meta VALUES ('schema', '1')")
        with self.assertRaisesRegex(ValueError, 'rebuild it: powo-nomenclature index'):
            Checklist(old)

    def test_cli_prints_json(self):
        out = io.StringIO()
        with redirect_stdout(out):
            main(['name', 'Veronica micrantha', '--index', str(self.index), '--author', 'Hoff. et LK.'])
        data = json.loads(out.getvalue())
        self.assertEqual((data['method'], data['fields']['powo_author']), ('exact', 'Hoffmanns. & Link'))


@unittest.skipUnless(__import__('importlib').util.find_spec('openpyxl'), 'openpyxl not installed (excel extra)')
class ExcelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.index = Path(cls.tmp.name) / 'wcvp.sqlite'
        build_index(FIXTURE, cls.index)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_workbook_fills_only_empty_fields_and_keeps_review_rows_out_of_the_import(self):
        from openpyxl import Workbook, load_workbook
        from powo_nomenclature.excel import process_workbook
        source = Path(self.tmp.name) / 'export.xlsx'
        book = Workbook()
        book.active.append(['id', 'revision', 'porto_genus', 'porto_specific_epithet', 'porto_author', 'porto_infra_rank',
                            'porto_infra_epithet', 'porto_infra_author', 'powo_family', 'powo_native'])
        book.active.append(['a', 3, 'Veronica', 'micrantha', 'Hoff. et LK.', '', '', '', 'Scrophulariaceae', ''])
        book.active.append(['b', 1, 'Salvia', 'officinalis', 'L.', 'proles', 'S. iberica', '', '', ''])
        book.active.append(['c', 2, '', '', '', '', '', '', '', ''])
        book.save(source)
        with Checklist(self.index) as checklist:
            review, imported, methods = process_workbook(source, checklist, Path(self.tmp.name) / 'out', lang='pt')
        self.assertEqual(sum(methods.values()), 2)
        rows = list(load_workbook(imported).active.iter_rows(values_only=True))
        self.assertEqual(rows[0][:3], ('id', 'revision', 'powo_family'))
        self.assertEqual([r[0] for r in rows[1:]], ['a'])  # «b» needs review.
        values = dict(zip(rows[0], rows[1]))
        self.assertEqual((values['revision'], values['powo_family'], values['powo_native']), (3, 'Scrophulariaceae', 'Portugal'))
        cell = load_workbook(imported).active['C2']
        self.assertEqual(cell.data_type, 's')
        sheets = load_workbook(review)
        self.assertEqual(sheets.sheetnames, ['Resultados', 'Diferenças', 'Fonte'])
        differences = list(sheets['Diferenças'].iter_rows(values_only=True))
        self.assertEqual([(r[3], r[4], r[5]) for r in differences[1:]], [('powo_family', 'Scrophulariaceae', 'Plantaginaceae')])

    def test_verify_workbook_lists_issues_per_row(self):
        from openpyxl import Workbook, load_workbook
        from powo_nomenclature.excel import verify_workbook
        source = Path(self.tmp.name) / 'filled.xlsx'
        book = Workbook()
        book.active.append(['id', 'porto_genus', 'porto_specific_epithet', 'powo_genus', 'powo_specific_epithet', 'powo_url'])
        book.active.append(['ok', 'Veronica', 'micrantha', 'Veronica', 'micrantha', url('812305-1')])
        book.active.append(['bad', 'Veronica', 'micrantha', 'Veronica', 'micrantha', url('99999999-9')])
        book.active.append(['empty', 'Veronica', 'micrantha', '', '', ''])
        book.save(source)
        with Checklist(self.index) as checklist:
            path, codes, checked = verify_workbook(source, checklist, Path(self.tmp.name) / 'out', lang='pt')
        self.assertEqual((checked, dict(codes)), (2, {'unknown_page': 1}))
        rows = list(load_workbook(path)['Problemas'].iter_rows(values_only=True))
        self.assertEqual([(r[1], r[3]) for r in rows[1:]], [('bad', 'unknown_page')])


if __name__ == '__main__':
    unittest.main()
